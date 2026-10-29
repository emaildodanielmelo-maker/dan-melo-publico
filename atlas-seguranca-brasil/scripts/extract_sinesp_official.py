"""Extrai uma série SINESP-VDE *independente* das tabelas FBSP.

São preservados os eventos publicados pelo Ministério da Justiça, sem juntar
Homicídio doloso e Feminicídio nem construir as categorias próprias do FBSP.

Fonte: https://www.gov.br/mj/pt-br/assuntos/sua-seguranca/seguranca-publica/estatistica/download/dnsp-base-de-dados
Entradas: /tmp/atlas-vde-{ano}.xlsx, de 2015 a 2026.
"""
from __future__ import annotations

import json
import math
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
STATES = {s['uf'] for s in json.loads((ROOT/'dist/dados.json').read_text())['states']}
YEARS = list(range(2015, 2027))

# Identificadores próprios; nomes correspondem às naturezas da planilha MJSP.
# Tipos selecionam a coluna quantitativa da própria fonte (sem misturar unidades).
EVENTS = [
 ('homicidio','Homicídio doloso','Letalidade','total_vitima','vítimas'),
 ('feminicidio','Feminicídio','Letalidade','total_vitima','vítimas'),
 ('tentativa_feminicidio','Tentativa de feminicídio','Letalidade','total_vitima','vítimas'),
 ('latrocinio','Roubo seguido de morte (latrocínio)','Letalidade','total_vitima','vítimas'),
 ('lesao_morte','Lesão corporal seguida de morte','Letalidade','total_vitima','vítimas'),
 ('intervencao','Morte por intervenção de Agente do Estado','Letalidade','total_vitima','vítimas'),
 ('tentativa_homicidio','Tentativa de homicídio','Letalidade','total_vitima','vítimas'),
 ('morte_agente','Morte de Agente do Estado','Outras ocorrências','total_vitima','vítimas'),
 ('morte_transito','Morte no trânsito ou em decorrência dele (exceto homicídio doloso)','Outras ocorrências','total_vitima','vítimas'),
 ('morte_esclarecer','Mortes a esclarecer (sem indício de crime)','Outras ocorrências','total_vitima','vítimas'),
 ('suicidio','Suicídio','Outras ocorrências','total_vitima','vítimas'),
 ('suicidio_agente','Suicídio de Agente do Estado','Outras ocorrências','total_vitima','vítimas'),
 ('estupro','Estupro','Violência sexual','total_vitima','vítimas'),
 ('estupro_vulneravel','Estupro de vulnerável','Violência sexual','total_vitima','vítimas'),
 ('roubo_veiculo','Roubo de veículo','Crimes patrimoniais','total','ocorrências'),
 ('furto_veiculo','Furto de veículo','Crimes patrimoniais','total','ocorrências'),
 ('roubo_carga','Roubo de carga','Crimes patrimoniais','total','ocorrências'),
 ('roubo_banco','Roubo a instituição financeira','Crimes patrimoniais','total','ocorrências'),
 ('trafico','Tráfico de drogas','Drogas e apreensões','total','ocorrências'),
 ('cocaina','Apreensão de Cocaína','Drogas e apreensões','total_peso','kg apreendidos'),
 ('maconha','Apreensão de Maconha','Drogas e apreensões','total_peso','kg apreendidos'),
 ('armas','Arma de Fogo Apreendida','Drogas e apreensões','total','armas apreendidas'),
 ('desaparecidos','Pessoa Desaparecida','Pessoas desaparecidas','total_vitima','pessoas'),
 ('localizados','Pessoa Localizada','Pessoas desaparecidas','total_vitima','pessoas'),
 ('mandados','Mandado de prisão cumprido','Atividade operacional','total','mandados'),
 ('atendimento','Atendimento pré-hospitalar','Atividade operacional','total','atendimentos'),
 ('salvamento','Busca e salvamento','Atividade operacional','total','atendimentos'),
 ('incendios','Combate a incêndios','Atividade operacional','total','atendimentos'),
 ('alvaras','Emissão de Alvarás de licença','Atividade operacional','total','alvarás'),
 ('vistorias','Realização de vistorias','Atividade operacional','total','vistorias'),
]


def extract():
    meta = [{'id':slug,'name':name,'theme':theme,'field':field,'numberType':unit,
             'denominator':'mulheres' if slug in ('feminicidio','tentativa_feminicidio') else 'habitantes','event':name}
            for slug,name,theme,field,unit in EVENTS]
    by_name = {name: (slug,field) for slug,name,_,field,_ in EVENTS}
    data = {m['id']:{uf:{} for uf in STATES} for m in meta}
    for year in YEARS:
        file = Path(f'/tmp/atlas-vde-{year}.xlsx')
        if not file.exists():
            raise FileNotFoundError(f'Arquivo oficial ausente: {file}')
        wb = openpyxl.load_workbook(file,read_only=True,data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = next(it)
        cols = {k:header.index(k) for k in ('evento','uf','data_referencia','abrangencia','total','total_vitima','total_peso')}
        sums = defaultdict(float)
        reported = set()
        for i,row in enumerate(it,2):
            if i%200000==0:
                print(f'{year}: {i:,} / {ws.max_row:,} linhas',flush=True)
            if row[cols['abrangencia']]!='Estadual':
                continue
            event,uf,dt = row[cols['evento']],row[cols['uf']],row[cols['data_referencia']]
            if event not in by_name or uf not in STATES or dt is None or dt.year!=year:
                continue
            slug,field=by_name[event]
            value=row[cols[field]]
            if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value) or value<0:
                continue
            key=(slug,uf,dt.month)
            sums[key]+=value
            reported.add(key)
        wb.close()
        for slug in data:
            for uf in STATES:
                data[slug][uf][str(year)] = [round(sums[(slug,uf,month)],3) if (slug,uf,month) in reported else None for month in range(1,13)]
        print(f'{year}: {len(reported):,} combinações mês/UF/evento informadas',flush=True)
    out={'source':'SINESP-VDE — MJSP (série nativa independente)','sourceUrl':'https://www.gov.br/mj/pt-br/assuntos/sua-seguranca/seguranca-publica/estatistica/download/dnsp-base-de-dados',
         'sourceFileUpdated':'2026-10-19','years':YEARS,'max2026Month':9,'metrics':meta,'data':data,
         'notes':['Categorias originais do SINESP-VDE: homicídio doloso exclui feminicídio; nunca soma entre si fontes criminais FBSP e MJSP.',
                  'A lista inclui indicadores criminais e operacionais. Não somar categorias com unidades ou fatos diferentes.',
                  'Usamos somente abrangência Estadual; células ausentes são nulas e zeros são mantidos como informados.',
                  '2026 disponível de janeiro a setembro e sujeito a retificações; YTD compara os mesmos meses de todos os anos.']}
    assert len(out['metrics'])==30
    path=ROOT/'dist/sinesp-oficial.json'
    path.write_text(json.dumps(out,ensure_ascii=False,separators=(',',':')),'utf-8')
    print(f'Gerado {path}: {path.stat().st_size:,} bytes',flush=True)


if __name__=='__main__':
    extract()
