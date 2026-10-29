"""Consolida meses comparáveis do SINESP-VDE 2021–2026.

Insumos oficiais: BancoVDE {ano}.xlsx, disponibilizados pelo MJSP em
https://www.gov.br/mj/pt-br/assuntos/sua-seguranca/seguranca-publica/estatistica/download/dnsp-base-de-dados

Salve cada arquivo em /tmp/atlas-vde-{ano}.xlsx e execute este script.
O histórico anual do Atlas continua sendo exclusivamente o FBSP. Esta base
forma uma SEGUNDA série, de metodologia SINESP, para comparações YTD.
"""
from __future__ import annotations

import json
import math
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
CURRENT = json.loads((ROOT / 'dist/dados.json').read_text(encoding='utf-8'))
STATES = {s['uf'] for s in CURRENT['states']}
OUTPUT = ROOT / 'dist/ytd-sinesp.json'
YEARS = list(range(2021, 2027))

# Correspondência com o Atlas apenas quando a definição da natureza é
# reconhecível no MJSP. Mesmo nomes similares não devem mesclar FBSP/SINESP.
# Para 2026, a fonte permite apenas o recorte Jan–Set divulgado em outubro.
DIRECT = {
    'homicidio': ('Homicídio doloso', 'total_vitima'),
    'homicidio-mulheres': ('Homicídio doloso', 'feminino'),
    'latrocinio': ('Roubo seguido de morte (latrocínio)', 'total_vitima'),
    'intervencao-policial': ('Morte por intervenção de Agente do Estado', 'total_vitima'),
    'tentativa-homicidio': ('Tentativa de homicídio', 'total_vitima'),
    'roubo-carga': ('Roubo de carga', 'total'),
    'roubo-veiculo': ('Roubo de veículo', 'total'),
    'furto-veiculo': ('Furto de veículo', 'total'),
    'estupro': ('Estupro', 'total_vitima'),
    'estupro-vulneravel': ('Estupro de vulnerável', 'total_vitima'),
    'feminicidio': ('Feminicídio', 'total_vitima'),
    'desaparecidos': ('Pessoa Desaparecida', 'total_vitima'),
    '_lesao_seguida_morte': ('Lesão corporal seguida de morte', 'total_vitima'),
}
DERIVED = {
    'roubo-furto-veiculo': ('roubo-veiculo', 'furto-veiculo'),
    'estupro-total': ('estupro', 'estupro-vulneravel'),
    'mvi': ('homicidio', 'latrocinio', '_lesao_seguida_morte', 'intervencao-policial'),
}
BY_EVENT = defaultdict(list)
for metric, (event, field) in DIRECT.items():
    BY_EVENT[event].append((metric, field))


def number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(value) and value >= 0:
        return int(round(value))
    return None


def source_rows(year):
    filename = Path(f'/tmp/atlas-vde-{year}.xlsx')
    if not filename.is_file():
        raise FileNotFoundError(f'Falta baixar a planilha oficial: {filename}')
    wb = openpyxl.load_workbook(filename, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    head = next(rows)
    indexes = {field: head.index(field) for field in
               ('uf', 'evento', 'data_referencia', 'abrangencia', 'total_vitima', 'feminino', 'total')}
    assert ws.title == str(year), (year, ws.title)
    return wb, rows, indexes, ws.max_row


def extract():
    all_data = {metric: {uf: {} for uf in STATES} for metric in [*DIRECT, *DERIVED]}
    for year in YEARS:
        wb, rows, col, nrows = source_rows(year)
        values = defaultdict(int)
        reported = set()
        months = set()
        for nr, row in enumerate(rows, 2):
            if nr % 200000 == 0:
                print(f'{year}: {nr:,}/{nrows:,} linhas processadas', flush=True)
            if row[col['abrangencia']] != 'Estadual':
                continue  # PF/PRF são fontes separadas, não somar às estaduais.
            event, uf, dt = row[col['evento']], row[col['uf']], row[col['data_referencia']]
            if event not in BY_EVENT or uf not in STATES or not dt or dt.year != year:
                continue
            month = dt.month
            months.add(month)
            for metric, field in BY_EVENT[event]:
                value = number(row[col[field]])
                if value is None:
                    continue  # sem informação não equivale a zero
                key = (metric, uf, month)
                values[key] += value
                reported.add(key)
        wb.close()
        assert months, year
        for uf in STATES:
            for mid in DIRECT:
                all_data[mid][uf][str(year)] = [
                    values[(mid, uf, month)] if (mid, uf, month) in reported else None
                    for month in range(1, 13)
                ]
        for derived, components in DERIVED.items():
            for uf in STATES:
                components_rows = [all_data[mid][uf][str(year)] for mid in components]
                all_data[derived][uf][str(year)] = [
                    sum(vals) if all(v is not None for v in vals) else None
                    for vals in zip(*components_rows)
                ]
        # O SINESP registra Feminicídio fora da natureza Homicídio doloso.
        # O indicador agregado de homicídios do FBSP inclui feminicídios.
        # Base documental: notas metodológicas dos gestores estaduais do MJSP,
        # seção I – Homicídio doloso (explicitamente "exceto Feminicídio").
        # Sem a composição, a queda aparente de homicídios seria superestimada.
        for uf in STATES:
            female_fem = all_data['feminicidio'][uf][str(year)]
            for mid in ('homicidio', 'homicidio-mulheres', 'mvi'):
                original = all_data[mid][uf][str(year)]
                all_data[mid][uf][str(year)] = [
                    (a+b if a is not None and b is not None else None)
                    for a, b in zip(original, female_fem)
                ]
        print(f'{year}: meses de referência {sorted(months)}; UFs {len(STATES)}', flush=True)

    output = {
        'source': 'SINESP-VDE / Ministério da Justiça e Segurança Pública',
        'sourceUrl': 'https://www.gov.br/mj/pt-br/assuntos/sua-seguranca/seguranca-publica/estatistica/download/dnsp-base-de-dados',
        'sourceFileUpdated': '2026-10-19',
        'years': YEARS,
        'max2026Month': max(i for i in range(1, 13) if any(
            vals.get('2026', [None]*12)[i-1] is not None for geo in all_data.values() for vals in geo.values()
        )),
        'mapping': {mid: {'type': 'direto', 'event': event, 'field': field} for mid, (event, field) in DIRECT.items() if not mid.startswith('_')},
        'data': {mid: by_uf for mid, by_uf in all_data.items() if not mid.startswith('_')},
        'notes': [
            '2021–2026 são uma série mensal independente do SINESP, não os totais anuais do FBSP.',
            'No SINESP, Homicídio doloso exclui Feminicídio; somamos Feminicídio às categorias agregadas de homicídio, homicídio de mulheres e MVI.',
            'Feminicídio e mortes por intervenção policial usam registros estaduais; não somar PF/PRF duas vezes.',
            'MVI é uma aproximação calculada como homicídio doloso + feminicídio + latrocínio + lesão corporal seguida de morte + morte por intervenção.',
            'Algumas UFs registram zeros estruturais ou revisam dados. Cobertura de linhas não garante qualidade do registro.'
        ],
    }
    for mid, comps in DERIVED.items():
        output['mapping'][mid] = {'type': 'combinado', 'components': [x.replace('_lesao_seguida_morte', 'lesão corporal seguida de morte') for x in comps]}
    output['mapping']['homicidio'] = {'type': 'combinado', 'components': ['Homicídio doloso (SINESP, sem feminicídio)', 'Feminicídio']}
    output['mapping']['homicidio-mulheres'] = {'type': 'combinado', 'components': ['Homicídio doloso · vítimas femininas (SINESP)', 'Feminicídio']}
    output['mapping']['mvi'] = {'type': 'combinado', 'components': ['Homicídio doloso (incluindo feminicídio)', 'Latrocínio', 'Lesão corporal seguida de morte', 'Morte por intervenção de agente do Estado']}
    assert len(output['data']) == 15
    assert output['max2026Month'] in (8, 9, 10), output['max2026Month']
    OUTPUT.write_text(json.dumps(output, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'Salvo {OUTPUT}, {OUTPUT.stat().st_size:,} bytes.', flush=True)
    for mid in output['data']:
        print(mid, 'meses2026', [sum(output['data'][mid][uf]['2026'][m] is not None for uf in STATES) for m in range(output['max2026Month'])])


if __name__ == '__main__':
    extract()
